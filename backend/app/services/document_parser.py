"""不可信文档的资源受限解析。

格式解析库运行在可终止的子进程中；压缩包规模、PDF 页数、提取字符数、并发数和执行时间
分别设限，降低压缩炸弹、恶意文件和 CPU/内存耗尽对 API 主进程的影响。
"""

import io
import multiprocessing
import threading
import zipfile
from html.parser import HTMLParser
from pathlib import Path


SUPPORTED_EXTENSIONS = {".md", ".txt", ".docx", ".pptx", ".pdf", ".html", ".htm"}
MAX_ARCHIVE_ENTRIES = 2_000
MAX_ARCHIVE_ENTRY_SIZE = 20 * 1024 * 1024
MAX_ARCHIVE_TOTAL_SIZE = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_PDF_PAGES = 200
MAX_EXTRACTED_CHARACTERS = 2_000_000
_PARSE_SLOTS = threading.BoundedSemaphore(2)


def _validate_ooxml_archive(data: bytes) -> None:
    """在交给 Office 解析库前检查 ZIP 条目数、展开大小和压缩比。"""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            total_size = sum(entry.file_size for entry in entries)
            unsafe = (
                len(entries) > MAX_ARCHIVE_ENTRIES
                or total_size > MAX_ARCHIVE_TOTAL_SIZE
                or any(entry.file_size > MAX_ARCHIVE_ENTRY_SIZE for entry in entries)
                or any(
                    entry.file_size > 0
                    and (entry.compress_size == 0 or entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO)
                    for entry in entries
                )
            )
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 Office 文档") from exc
    if unsafe:
        raise ValueError("Office 文档超过压缩资源限制")


def _join_bounded(lines) -> str:
    """边拼接边统计输出，避免解析完成后才发现文本已耗尽内存。"""
    result: list[str] = []
    length = 0
    for line in lines:
        length += len(line)
        if length > MAX_EXTRACTED_CHARACTERS:
            raise ValueError("文档提取内容过大")
        result.append(line)
    return "\n".join(result)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    _validate_ooxml_archive(data)
    doc = Document(io.BytesIO(data))
    return _join_bounded(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    _validate_ooxml_archive(data)
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > 500:
        raise ValueError("演示文稿页数过多")
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        lines.append(f"\n## 幻灯片 {index}\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
    return _join_bounded(lines)


def _parse_pdf(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF 页数过多")
        return _join_bounded(filter(None, (page.extract_text() for page in pdf.pages)))


class _ReadableHTMLParser(HTMLParser):
    ignored_tags = {"script", "style", "nav", "footer", "header"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.ignored_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.ignored_tags:
            self.ignored_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self.ignored_tags and self.ignored_depth:
            self.ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text and not self.ignored_depth:
            self.parts.append(text)


def _parse_html(data: bytes) -> str:
    parser = _ReadableHTMLParser()
    parser.feed(data.decode("utf-8", errors="replace"))
    return "\n".join(parser.parts)


def parse_document(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"不支持的格式: {extension}。支持: {supported}")
    if extension in {".md", ".txt"}:
        return data.decode("utf-8", errors="replace")
    if extension == ".docx":
        return _parse_docx(data)
    if extension == ".pptx":
        return _parse_pptx(data)
    if extension == ".pdf":
        return _parse_pdf(data)
    return _parse_html(data)


def _parse_worker(connection, filename: str, data: bytes) -> None:
    try:
        connection.send(("ok", parse_document(filename, data)))
    except Exception as exc:
        connection.send(("error", str(exc)))
    finally:
        connection.close()


def _terminate_process(process) -> None:
    """正常回收解析进程；超时后先 terminate，仍未退出再强制 kill。"""
    if not process.is_alive():
        process.join()
        return
    process.terminate()
    process.join(timeout=2)
    if process.is_alive():
        process.kill()
        process.join()


def parse_document_isolated(filename: str, data: bytes, timeout: float = 20) -> str:
    """在受并发限制、可超时终止的独立进程中解析不可信文档。"""
    with _PARSE_SLOTS:
        context = multiprocessing.get_context("spawn")
        receiver, sender = context.Pipe(duplex=False)
        process = context.Process(target=_parse_worker, args=(sender, filename, data), daemon=True)
        process.start()
        sender.close()
        try:
            if not receiver.poll(timeout):
                raise TimeoutError("文档解析超时")
            try:
                status, payload = receiver.recv()
            except EOFError as exc:
                raise ValueError("文档解析进程异常退出") from exc
        finally:
            receiver.close()
            _terminate_process(process)
    if status == "error":
        raise ValueError(payload)
    return payload
